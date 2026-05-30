"""
LangChain 版本的 RAG 核心逻辑

作用：将原有的 RAG 逻辑完全重构为 LangChain 版本，
      使用 LangChain 组件实现所有功能
      
关键知识点：
    1.组件化思想：将复杂功能拆分为可重用的组件
    2.Document 对象：LangChain 的标准文档格式
    3.Pipeline 编排：组件之间的数据流和调用关系
    4.错误处理：确保整个流程的健壮性
    
对比学习点：
- 如何使用更少的代码实现相同的功能
- LangChain 各组件如何协作
- 流程是如何通过 Runnable 编排的
"""
import pdfplumber
from langchain_core.documents import Document
from langchain_core.runnables import RunnablePassthrough
from langchain_core.output_parsers import StrOutputParser
from operator import itemgetter
from docx import Document as WordDocument
from openpyxl import load_workbook
from typing import List, Any, Dict
from utils.langchain_adapters import get_text_splitter,get_vectorstore,get_retriever,get_prompt_template,get_compression_retriever
from utils.logger import logger
from utils.config import settings
from utils.llm_chain import qwen_llm


# ==========================
# 1. 文件解析
# ==========================
# 解析纯文本文件（txt,md）
def extract_txt_md(file_path: str) -> str:
    """读取文本文件，自动兼容 utf-8/gbk 编码"""
    try:
        # 先尝试 utf-8
        with open(file_path, "r", encoding="utf-8") as f:
            return f.read()
    except UnicodeDecodeError:
        # 失败则用 gbk 编码（Windows 常用）
        with open(file_path, "r", encoding="gbk") as f:
            return f.read()

# 解析 pdf 文件
def extract_pdf(file_path: str) -> str:
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text

#  解析docx文件
def extract_docx(file_path:str) -> str:
    doc = WordDocument(file_path)
    return "\n".join([p.text.strip() for p in doc.paragraphs if p.text.strip()])

# 解析 xlsx 文件
def extract_xlsx(file_path: str) -> str:
    full_text = []
    # 读取xisx文件中的多个表格
    wb = load_workbook(filename=file_path, read_only=True, data_only=True)
    # 遍历获取每个表格
    for sheet in wb.worksheets:
        # 遍历获取每个表格中的每个单元格
        for row in sheet.iter_rows(values_only=True):
            # 将单元格中的内容存入变量
            row_data = [str(cell).strip() for cell in row if cell is not None]
            if row_data:
                full_text.append(" | ".join(row_data))
    return "\n".join(full_text)

# 路由函数：自动识别文件后缀 ——> 调用对应解析器
def get_file_content(file_path: str) -> str:
    ext = file_path.split(".")[-1].lower()
    if ext in ["txt", "md"]:
        return extract_txt_md(file_path)
    elif ext == "pdf":
        return extract_pdf(file_path)
    elif ext == "docx":
        return extract_docx(file_path)
    elif ext == "xlsx":
        return extract_xlsx(file_path)
    else:
        raise ValueError(f"不支持的文件类型：{ext}")
            

# ==========================
# 2. 文本分块
# ==========================
def split_document(file_path: str, source_name: str) -> List[Document]:
    """
    文本分块 - 使用 LangChain TextSplitter
    
    与原代码 split_text() 对比：
    - 原代码：返回纯字符串列表
    - 新版本：返回 Document 列表，每个 Document 自带 metadata
    - metadata 包含了文件的 source 信息，这对后续追踪答案来源非常重要
    """
    content = get_file_content(file_path)
    # 创建 Document 对象（带元数据）
    docs = [Document(page_content = content, metadata={"source": source_name})]
    # 使用 TextSplitter 分块
    splitter = get_text_splitter()
    splitted = splitter.split_documents(docs)
    logger.info(f"✅ 分块完成：{len(splitted)} 个块")
    return splitted

# ==========================
# 3. 向量入库
# ==========================
def add_to_vectorstore(documents: List[Document], collection_name: str = "knowledge_base"):
    """
    添加文档到向量库
    
    与原代码 add_chunks_to_chroma() 对比：
    - 原代码：需要手动调用 embedding API + 组装 IDs + 调用 chromadb.add()
    - 新版本：VectorStore 统一管理，自动调用 embedding + 批量添加
    """
    store = get_vectorstore(collection_name)
    ids = store.add_documents(documents)
    logger.info(f"✅ 添加 {len(documents)} 个文档到向量库")
    return ids

# ==========================
# 4. 检索文档
# ==========================
def retrieve_documents(query: str, top_k: int = None, use_rerank: bool = True) -> List[Dict]:
    """
    检索相关文档 - 使用 LangChain Retriever
    
    与原代码 retrieve_relevant_docs() 对比：
    - 原代码：手动调用 embedding API + Chroma.query() + 结果解析
    - 新版本：Retriever 一行代码搞定全部流程！
    
    Retriever 的优势：
    1. 代码更少，更易读
    2. 支持多种搜索策略（similarity, mmr, etc.）
    3. 易于扩展（添加 compression, filtering 等）
    """
    top_k =top_k or settings.recall_top_k
    
    if use_rerank:
        retriever = get_compression_retriever(top_k=top_k*2)
    else:
        retriever = get_retriever(top_k=top_k)
    
    # invoke() 是 LangChain Runnable 协议的标准调用方法
    # 执行检索：问题 → 向量 → 搜索 → 返回文档
    docs = retriever.invoke(query)
    
    # Document 对象已经包含 page_content 和 metadata
    # 把 LangChain 文档 → 转为干净的字典格式
    results = [
        {
            "content" : doc.page_content.strip(),
            "source" : doc.metadata.get("source","未知") 
        }
        for doc in docs
    ]
    logger.info(f"✅ 检索完成（{'重排' if use_rerank else '普通'}）：{len(results)} 条结果")
    return results
    

# =========================================
# 构建 Prompt
# =========================================

def build_rag_context(query: str, history: List[Dict] = None, top_k: int = None, use_rerank: bool = False) -> Dict[str,Any]:
    """
    构建 RAG 上下文
    
    这是一个中间函数，将检索结果和历史对话组织成
    LangChain Prompt 需要的输入格式
    """
    history = history or []
    
    # 格式化历史对话为 LangChain 消息格式
    chat_history = []
    for msg in history[-settings.max_chat_history_round:]:
        role = "assistant" if msg["role"] == "assistant" else "user"
        chat_history.append({"role": role, "content": msg["content"]})
    
    # 检索文档    
    retrieved = retrieve_documents(query, top_k, use_rerank)
    
    # 格式化上下文
    context = "\n\n".join(
        f"[来源：{doc['source']}]\n{doc['content']}"
        for doc in retrieved
    )
    
    return {
        "chat_history": chat_history,
        "context" : context,
        "question" : query
    }
    
# =========================================
# 完整的 LangChain RAG Pipeline
# =========================================
def create_full_rag_chain(
    use_rerank: bool = True,
    top_k: int = None
):
    """
    创建完整的 LangChain RAG 链路
    流程：问题 —> 检索 —> 重排 —> 提示词 —> LLM —> 回答
    """
    
    top_k = top_k or settings.recall_top_k
    
    # 创建检索器
    if use_rerank:
        retriever = get_compression_retriever(top_k=top_k*2)
    else:
        retriever = get_retriever(top_k=top_k)
        
    # 提示词和LLM
    prompt = get_prompt_template()
    llm = qwen_llm
    
    # 构建 RAG 管道
    # 构建输入字典
    rag_pipeline = (
        {
        "context": itemgetter("question") | retriever,  # 从输入 dict 中取出纯问题字符串 → 传给 retriever
        "question": itemgetter("question"),             # 取出问题字符串 → 传入 Prompt 的 {question}
        "chat_history": itemgetter("chat_history"),     # 取出历史列表 → 传入 Prompt 的 MessagesPlaceholder
        }
        | prompt
        | llm
        | StrOutputParser() #输出解析器：把 LLM 返回的复杂消息体，只提取纯文本
    )
    return rag_pipeline

# 便捷调用函数
def ask_with_full_chain(question: str) -> str:
    chain = create_full_rag_chain()
    return chain.invoke(question)
 